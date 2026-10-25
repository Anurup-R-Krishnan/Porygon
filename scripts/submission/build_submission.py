"""Build Anurup's individual technical-enrichment evidence pack from local artifacts."""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "artifacts/local/college-submission"
OUT = WORK / "anurup-technical-enrichment-submission"
PUBLIC = "https://github.com/Anurup-R-Krishnan/Porygon/tree/codex/technical-enrichment"
DATASET_URL = PUBLIC + "/datasets/porygon-pilot-v1.0.0"


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    cert = OUT / "01-professional-online-certification"
    cert.mkdir(exist_ok=True)
    for source, target in [
        ("01_Course_Details_Anurup_R_Krishnan.pdf", "course-details.pdf"),
        (
            "02_One_Page_Project_Relevance_Summary.pdf",
            "one-page-project-relevance-summary.pdf",
        ),
        (
            "03_Original_Cisco_Completion_Certificate.pdf",
            "original-cisco-completion-certificate.pdf",
        ),
    ]:
        shutil.copy2(WORK / "cisco-source" / source, cert / target)
    assert len(PdfReader(cert / "one-page-project-relevance-summary.pdf").pages) == 1
    dataset = OUT / "02-dataset-release"
    dataset.mkdir(exist_ok=True)
    shutil.copytree(
        ROOT / "datasets/porygon-pilot-v1.0.0",
        dataset / "porygon-pilot-v1.0.0",
        dirs_exist_ok=True,
        ignore=shutil.ignore_patterns("__pycache__"),
    )
    repo = OUT / "03-open-source-technical-repository"
    repo.mkdir(exist_ok=True)
    history = subprocess.check_output(
        ["git", "log", "-30", "--date=short", "--format=%h | %ad | %an | %s"],
        cwd=ROOT,
        text=True,
    )
    (repo / "commit-history.txt").write_text(history)
    sha = subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
    ).strip()
    (repo / "repository-details.json").write_text(
        json.dumps(
            {
                "public_repository": PUBLIC,
                "local_source_commit": sha,
                "working_tree_changes_included": False,
                "public_dataset_release_url": DATASET_URL,
                "prepared_for": "Anurup R Krishnan",
            },
            indent=2,
        )
        + "\n"
    )
    shutil.copy2(
        ROOT / "docs/submission/installation-guide.md", repo / "installation-guide.md"
    )
    subprocess.run(
        [
            "git",
            "archive",
            "--format=zip",
            "--output=" + str(repo / "source-snapshot.zip"),
            "HEAD",
        ],
        cwd=ROOT,
        check=True,
    )
    shutil.copy2(ROOT / "README.md", repo / "project-readme.md")
    shutil.copy2(ROOT / "compose.yaml", repo / "compose.yaml")
    shutil.copy2(ROOT / "Makefile", repo / "Makefile")
    demo = OUT / "04-working-demonstration"
    demo.mkdir(exist_ok=True)
    recording = WORK / "recording"
    for name in [
        "demo.mp4",
        "demo.cast",
        "run.log",
        "run-metadata.json",
        "metrics.json",
        "versions.json",
        "narration.txt",
    ]:
        if (recording / name).exists():
            shutil.copy2(recording / name, demo / name)
    shutil.copy2(
        ROOT / "docs/submission/demonstration-instructions.md",
        demo / "demonstration-instructions.md",
    )
    for name in [
        "run_demo.sh",
        "record_demo.py",
        "render_demo.py",
        "collect_demo_metrics.py",
        "reproduce_metrics.py",
    ]:
        if (ROOT / "scripts/submission" / name).exists():
            shutil.copy2(ROOT / "scripts/submission" / name, demo / name)
    evidence = demo / "sample-outputs"
    evidence.mkdir(exist_ok=True)
    for pattern in [
        "ebpf-capture-integrity.json",
        "scoring-config.json",
        "scoring-*-score.json",
        "detection-ruleset.json",
        "detection-novel-detection.json",
        "containment-*.json",
    ]:
        for p in (recording / "sample-outputs").glob(pattern):
            shutil.copy2(p, evidence / p.name)
    ready = (
        (demo / "demo.mp4").exists()
        and (demo / "run-metadata.json").exists()
        and json.loads((demo / "run-metadata.json").read_text())["exit_code"] == 0
    )
    dataset_count = json.loads(
        (dataset / "porygon-pilot-v1.0.0/metadata.json").read_text()
    )["trial_count"]
    pages = [
        (
            "Technical Enrichment Activities",
            "Individual contribution submission",
            [
                ("p", "Anurup R Krishnan · cb.sc.u4cse23559"),
                (
                    "p",
                    "Amrita School of Computing, Amrita Vishwa Vidyapeetham, Coimbatore",
                ),
                (
                    "p",
                    "Project: Porygon — Container Runtime Anomaly Detection Research",
                ),
                ("p", "Course: 23CSE498 Project Phase II · Prepared 25 October 2026"),
                ("h", "Submission index"),
                (
                    "table",
                    [
                        ["PDF component", "Marks", "Evidence folder"],
                        [
                            "Professional Online Certification (individual)",
                            "20",
                            "01-professional-online-certification",
                        ],
                        [
                            "Dataset Release (activity 3, group)",
                            "5",
                            "02-dataset-release",
                        ],
                        [
                            "Open-source Technical Repository (activity 5, group)",
                            "5",
                            "03-open-source-technical-repository",
                        ],
                        [
                            "Working demonstration (activity 8, group)",
                            "5",
                            "04-working-demonstration",
                        ],
                    ],
                ),
                (
                    "p",
                    "This pack records Anurup’s certification and contributions to the three selected group activities. The college evaluates the group activities once per team.",
                ),
                ("h", "Format and routing"),
                (
                    "p",
                    "Organized in the component order and expected-deliverable categories of Technical Enrichment Activities.pdf. Submit supporting evidence through the project guide for verification and consolidated OneDrive upload.",
                ),
                (
                    "p",
                    "Deadline in the supplied document: one week before the last working day of semester seven. The calendar deadline is set by the department.",
                ),
            ],
        ),
        (
            "Professional Online Certification",
            "Mandatory individual activity · 20 marks",
            [
                (
                    "table",
                    [
                        ["Required item", "Submitted evidence"],
                        [
                            "Completion certificate",
                            "Original Cisco PDF, preserved without modification",
                        ],
                        ["Course details", "Course-details PDF"],
                        [
                            "One-page project relevance summary",
                            "Exactly one A4 page; attached",
                        ],
                    ],
                ),
                ("h", "Course details"),
                (
                    "p",
                    "Cyber Threat Management · Cisco Networking Academy · Online, self-paced · 16 hours (as recorded in the course-details document) · Completed 24 October 2026.",
                ),
                ("h", "Project application"),
                (
                    "p",
                    "The course supports vulnerability assessment, threat intelligence, risk-based controls, digital forensics and incident response in Porygon. The submitted one-page summary maps these topics to digest-bound telemetry, Trivy vulnerability context, incident evidence and operator-approved containment.",
                ),
                ("h", "Evidence identification"),
                (
                    "p",
                    "Student identity and register number follow the existing Cisco submission documents. The original completion certificate is included separately so its credential-verification details remain intact. Guide approval is recorded as student-confirmed in the existing course-details document.",
                ),
            ],
        ),
        (
            "Dataset Release",
            "Activity 3 · contribution to the group deliverable · 5 marks",
            [
                (
                    "table",
                    [
                        ["Expected deliverable", "Package item"],
                        ["Dataset documentation", "porygon-pilot-v1.0.0/readme.md"],
                        [
                            "Metadata and annotation guidelines",
                            "metadata.json, data-dictionary.md, annotation.md",
                        ],
                        [
                            "License information",
                            "license.txt and license.md — CC BY 4.0",
                        ],
                        [
                            "Version details",
                            "Version 1.0.0; source hashes and checksums",
                        ],
                        [
                            "Organized dataset repository",
                            "216 original trial JSON files, provenance and CSV index",
                        ],
                        ["Public repository link", DATASET_URL],
                    ],
                ),
                ("h", "Dataset description"),
                (
                    "p",
                    f"{dataset_count} completed controlled-container trials: three workloads (nginx, Redis, PostgreSQL), three scenarios, two context variants and twelve replicas per cell. The index contains 72 controlled non-attack-like trials and 144 controlled attack-like simulations. Evidence class: pilot. Granularity: complete-trial metadata.",
                ),
                ("h", "Reproducibility"),
                (
                    "code",
                    "python3 datasets/porygon-pilot-v1.0.0/verify.py\npython3 scripts/package_pilot_dataset.py",
                ),
                (
                    "p",
                    "The verifier checks file integrity, unique trial IDs, labels against source records and the 18-cell collection matrix. CC BY 4.0 attribution is included in the dataset license notice.",
                ),
            ],
        ),
        (
            "Open-source Technical Repository",
            "Activity 5 · contribution to the group deliverable · 5 marks",
            [
                ("p", "Public project repository: " + PUBLIC),
                (
                    "table",
                    [
                        ["Expected deliverable", "Evidence"],
                        [
                            "Source code",
                            "Public repository; source snapshot in this folder",
                        ],
                        [
                            "Well-structured organization",
                            "backend, collector, telemetry, responder, scanner, dashboard, experiments",
                        ],
                        [
                            "README and installation guide",
                            "project-readme.md and installation-guide.md",
                        ],
                        [
                            "Documentation",
                            "Architecture, research protocol and verification documentation in the repository",
                        ],
                        [
                            "Regular commits with version history",
                            "commit-history.txt; author names and commit identifiers retained",
                        ],
                    ],
                ),
                ("h", "Architecture and contribution evidence"),
                (
                    "p",
                    "Falco observes process execution; the telemetry adapter delivers events durably. The Docker collector supplies lifecycle and immutable image identity. The backend builds baselines, computes scores and produces incidents. The responder executes approved container actions. The scanner provides image vulnerability and SBOM context.",
                ),
                (
                    "p",
                    "The attached commit history identifies recorded authorship. This pack includes Anurup’s dataset packaging, verification workflow and submission automation.",
                ),
                ("h", "Version identification"),
                (
                    "p",
                    "Local source commit: "
                    + sha
                    + ". The source snapshot contains the committed repository at this revision. The public submission branch includes the dataset and reproduction instructions.",
                ),
            ],
        ),
        (
            "Working demonstration",
            "Activity 8 · contribution to the group deliverable · 5 marks",
            [
                (
                    "table",
                    [
                        ["Expected deliverable", "Evidence"],
                        [
                            "Working end-to-end prototype",
                            "Telemetry → baseline → scoring → incident → approved containment",
                        ],
                        [
                            "Short video",
                            "demo.mp4" if ready else "Recording in preparation",
                        ],
                        ["Instructions", "demonstration-instructions.md"],
                        [
                            "Commands/script regenerating metrics",
                            "run_demo.sh and reproduce_metrics.py",
                        ],
                        [
                            "Sample outputs and versions",
                            "sample-outputs/, metrics.json, versions.json",
                        ],
                    ],
                ),
                ("h", "Demonstration sequence"),
                (
                    "p",
                    "Start the stack and confirm service readiness. Execute a controlled workload, capture process events and bind the container to its image digest. Build and activate a baseline. Execute a novel process sequence, compute its score and create an explainable incident. Demonstrate operator credential separation, exact-target containment and rollback. Verify the dataset independently.",
                ),
                ("h", "Recorded result"),
                (
                    "p",
                    "The recorded end-to-end verification completed successfully; the accompanying log and sample outputs contain the acceptance assertions."
                    if ready
                    else "The live recording and final metrics will be attached after the end-to-end verification completes.",
                ),
                ("h", "Submission completion"),
                (
                    "p",
                    "The project guide verifies the activity selection and supporting files, then uploads the consolidated submission to the designated college location. Public repository and dataset links are included in the activity evidence.",
                ),
            ],
        ),
    ]
    for name, filename in [
        ("SubmissionSans", "DejaVuSans.ttf"),
        ("SubmissionBold", "DejaVuSans-Bold.ttf"),
        ("SubmissionMono", "DejaVuSansMono.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(name, "/usr/share/fonts/TTF/" + filename))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = (
            "SubmissionBold"
            if "Heading" in style.name or style.name == "Title"
            else "SubmissionSans"
        )

    styles.add(
        ParagraphStyle(
            name="BodySmall",
            fontName="SubmissionSans",
            fontSize=10,
            leading=15,
            spaceAfter=9,
            textColor=colors.HexColor("#243145"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="TableSmall", fontName="SubmissionSans", fontSize=9, leading=12
        )
    )
    styles.add(
        ParagraphStyle(
            name="CodeSmall",
            fontName="SubmissionMono",
            fontSize=8,
            leading=12,
            backColor=colors.HexColor("#eef2f6"),
            spaceAfter=10,
        )
    )
    story = []
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.core_properties.author = "Anurup R Krishnan"

    for i, (title, subtitle, blocks) in enumerate(pages):
        if i:
            story.append(PageBreak())
            doc.add_page_break()
        story.extend(
            [
                Paragraph(escape(title), styles["Title"]),
                Paragraph(escape(subtitle), styles["BodySmall"]),
                Spacer(1, 8 * mm),
            ]
        )
        doc.add_heading(title, 0)
        doc.add_paragraph(subtitle)
        for kind, value in blocks:
            if kind == "table":
                rows = [
                    [Paragraph(escape(str(c)), styles["TableSmall"]) for c in row]
                    for row in value
                ]
                widths = (
                    [52 * mm, 18 * mm, 100 * mm]
                    if len(value[0]) == 3
                    else [61 * mm, 109 * mm]
                )
                table = Table(rows, colWidths=widths, hAlign="LEFT")
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e4ebf2")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                            ("TOPPADDING", (0, 0), (-1, -1), 8),
                            (
                                "LINEBELOW",
                                (0, 0),
                                (-1, -1),
                                0.4,
                                colors.HexColor("#ccd5df"),
                            ),
                        ]
                    )
                )
                story.extend([table, Spacer(1, 5 * mm)])
                t = doc.add_table(rows=0, cols=len(value[0]))
                t.style = "Light Shading Accent 1"
                for row in value:
                    cells = t.add_row().cells
                    for j, c in enumerate(row):
                        cells[j].text = str(c)
            elif kind == "h":
                story.append(Paragraph(escape(value), styles["Heading2"]))
                doc.add_heading(value, 2)
            elif kind == "code":
                story.append(
                    Paragraph(escape(value).replace("\n", "<br/>"), styles["CodeSmall"])
                )
                doc.add_paragraph(value, style="No Spacing")
            else:
                story.append(Paragraph(escape(value), styles["BodySmall"]))
                doc.add_paragraph(value)

    def footer(canvas, document):
        canvas.setFont("SubmissionSans", 8)
        canvas.setFillColor(colors.HexColor("#526274"))
        canvas.drawString(
            20 * mm, 13 * mm, "PORYGON | ANURUP R KRISHNAN | TECHNICAL ENRICHMENT"
        )
        canvas.drawRightString(190 * mm, 13 * mm, str(document.page))

    pdf = OUT / "submission-evidence.pdf"
    SimpleDocTemplate(
        str(pdf),
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=22 * mm,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    doc.save(OUT / "submission-evidence.docx")
    writer = PdfWriter()
    for p in [
        pdf,
        cert / "course-details.pdf",
        cert / "one-page-project-relevance-summary.pdf",
        cert / "original-cisco-completion-certificate.pdf",
    ]:
        writer.append(p)
    writer.write(OUT / "consolidated-submission.pdf")
    (OUT / "readme.md").write_text(
        """# Anurup — technical enrichment submission\n\nStart with `consolidated-submission.pdf`. The editable evidence index is\n`submission-evidence.docx`. Original Cisco documents are in folder 01; the\nCC BY 4.0 dataset is in folder 02; repository evidence is in folder 03;\nthe video, reproducible commands and measured outputs are in folder 04.\n\nThe activity order follows Technical Enrichment Activities.pdf.\nOnly Anurup’s individual contribution is included.\n"""
    )
    hashes = {
        str(p.relative_to(OUT)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(OUT.rglob("*"))
        if p.is_file() and p.name != "submission-checksums.json"
    }
    (OUT / "submission-checksums.json").write_text(json.dumps(hashes, indent=2) + "\n")
    archive = WORK / "anurup-technical-enrichment-submission.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                z.write(p, Path(OUT.name) / p.relative_to(OUT))
    print(
        json.dumps(
            {
                "output": str(OUT),
                "pages": len(PdfReader(OUT / "consolidated-submission.pdf").pages),
                "video_included": ready,
                "zip": str(archive),
            }
        )
    )


if __name__ == "__main__":
    build()
